import os
import openpyxl
from datetime import datetime
import calendar
import random

def create_statement(filename, account_holder, months_count, start_year, start_month, base_salary, is_salaried, salary_employer, foir_pct, surplus_pct, has_bounces, variance_pct=0.0):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Statement"
    
    # Metadata
    ws["A1"] = "SYNTHETIC BANK - ACCOUNT STATEMENT"
    ws["A1"].font = openpyxl.styles.Font(bold=True, size=14)
    ws["A2"] = "Account Holder:"
    ws["B2"] = account_holder
    ws["A3"] = "Account Number:"
    ws["B3"] = f"SYN{random.randint(1000000, 9999999)}"
    ws["A4"] = "Bank Name:"
    ws["B4"] = "Synthetic Bank"
    ws["A5"] = "Statement Period:"
    
    # Calculate end date
    end_year, end_month = start_year, start_month
    for _ in range(months_count - 1):
        end_month += 1
        if end_month > 12:
            end_month = 1
            end_year += 1
            
    last_day = calendar.monthrange(end_year, end_month)[1]
    start_date_str = datetime(start_year, start_month, 1).strftime("%d %b %Y")
    end_date_str = datetime(end_year, end_month, last_day).strftime("%d %b %Y")
    ws["B5"] = f"{start_date_str} - {end_date_str}"
    
    ws["A6"] = "Opening Balance:"
    ws["B6"] = 50000.00
    
    # Headers
    headers = ["Date", "Description", "Debit", "Credit", "Balance"]
    for col_idx, text in enumerate(headers, 1):
        cell = ws.cell(row=8, column=col_idx)
        cell.value = text
        cell.font = openpyxl.styles.Font(bold=True)
        cell.fill = openpyxl.styles.PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        cell.font = openpyxl.styles.Font(color="FFFFFF", bold=True)
        
    current_balance = 50000.00
    row_num = 9
    
    curr_y, curr_m = start_year, start_month
    
    merchants = [
        "UPI-SWIGGY-FOOD", "UPI-ZOMATO-DINING", "CARD-AMAZON-RETAIL", 
        "UPI-UBER-TRAVEL", "CARD-RELIANCE FRESH", "UPI-BLINKIT", 
        "UPI-ZEPTO", "CARD-DMART", "UPI-PAYTM-QR", "UPI-PHONEPE-MERCHANT",
        "CARD-STARBUCKS", "UPI-LOCAL VENDOR", "UPI-APOLLO PHARMACY",
        "CARD-NETFLIX", "CARD-SPOTIFY"
    ]
    
    for month_idx in range(months_count):
        month_txs = []
        
        # Determine actual salary this month (apply variance if requested)
        actual_salary = base_salary
        if variance_pct > 0 and month_idx > 0:
            if random.random() > 0.6:
                drop = base_salary * (variance_pct / 100.0)
                actual_salary = base_salary - drop
        
        # Incomes
        if is_salaried:
            salary_day = 28 if curr_m == 2 else 30
            salary_desc = f"SALARY CREDIT - {salary_employer}"
            month_txs.append((datetime(curr_y, curr_m, salary_day), salary_desc, None, actual_salary))
        else:
            # Freelance/Gig income (3 to 6 gig payments)
            for _ in range(random.randint(3, 6)):
                gig_day = random.randint(1, 28)
                gig_amt = actual_salary * random.uniform(0.1, 0.3)
                month_txs.append((datetime(curr_y, curr_m, gig_day), f"UPI-CLIENT PAYMENT {random.randint(100,999)}", None, gig_amt))
                
        # Random small credits (friends paying back, cashback, etc.)
        for _ in range(random.randint(2, 5)):
            c_day = random.randint(1, 28)
            c_amt = random.randint(100, 2500)
            month_txs.append((datetime(curr_y, curr_m, c_day), f"UPI-CREDIT-FRIEND-{random.randint(1000,9999)}", None, c_amt))
                
        # Fixed Obligations (EMIs)
        emi_amount = base_salary * (foir_pct / 100.0)
        if emi_amount > 0:
            # Split into 1 to 3 EMIs depending on size
            emi_count = 1 if emi_amount < 15000 else random.randint(2, 3)
            split_emi = emi_amount / emi_count
            for e in range(emi_count):
                emi_day = 5 + (e * 5) # e.g. 5th, 10th, 15th
                month_txs.append((datetime(curr_y, curr_m, emi_day), f"ACH-LOAN EMI-{e}", split_emi, None))
            
        # Bounces
        if has_bounces and month_idx == (months_count // 2):
            month_txs.append((datetime(curr_y, curr_m, 4), "ACH-BAJAJ FIN EMI", 15000, None))
            month_txs.append((datetime(curr_y, curr_m, 4), "ACH RETURN-INSUFFICIENT FUNDS", None, 15000))
            month_txs.append((datetime(curr_y, curr_m, 5), "ACH BOUNCE PENALTY", 590, None))
            
        # Living Expenses (High Volume)
        living_exp = base_salary * (1.0 - (foir_pct/100.0) - (surplus_pct/100.0))
        
        # We want approx 40-70 transactions a month to look realistic
        num_transactions = random.randint(40, 70)
        
        # Distribute living expenses randomly
        remaining_exp = living_exp
        
        for i in range(num_transactions):
            if i == num_transactions - 1:
                chunk_amt = remaining_exp # Put all remaining in last transaction to balance exactly
            else:
                # Average chunk size
                avg_chunk = remaining_exp / (num_transactions - i)
                # Randomize between 20% and 180% of average to get small and big purchases
                chunk_amt = avg_chunk * random.uniform(0.2, 1.8)
                
            chunk_amt = round(chunk_amt, 2)
            if chunk_amt <= 0: continue
                
            remaining_exp -= chunk_amt
            day = random.randint(1, 28)
            month_txs.append((datetime(curr_y, curr_m, day), random.choice(merchants), chunk_amt, None))
            
        # Sort and write
        month_txs.sort(key=lambda x: x[0])
        for date_obj, desc, debit, credit in month_txs:
            ws.cell(row=row_num, column=1).value = date_obj.strftime("%d %b %Y")
            ws.cell(row=row_num, column=2).value = desc
            if debit is not None:
                ws.cell(row=row_num, column=3).value = round(debit, 2)
                current_balance -= debit
            else:
                ws.cell(row=row_num, column=3).value = ""
                
            if credit is not None:
                ws.cell(row=row_num, column=4).value = round(credit, 2)
                current_balance += credit
            else:
                ws.cell(row=row_num, column=4).value = ""
                
            ws.cell(row=row_num, column=5).value = round(current_balance, 2)
            row_num += 1
            
        # Next month
        curr_m += 1
        if curr_m > 12:
            curr_m = 1
            curr_y += 1
            
    # Auto-adjust column widths
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(cell.value)
            except:
                pass
        adjusted_width = (max_length + 2)
        ws.column_dimensions[column].width = adjusted_width

    wb.save(filename)
    print(f"Generated {filename}")

if __name__ == "__main__":
    out_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'test_datasets')
    if not os.path.exists(out_dir):
        os.makedirs(out_dir)
        
    print("Generating HIGH VOLUME synthetic datasets...")
    
    # 1. Gold Excellent
    create_statement(
        filename=os.path.join(out_dir, 'scenario_gold_excellent.xlsx'),
        account_holder="Gold User",
        months_count=12,
        start_year=2024,
        start_month=6,
        base_salary=250000,
        is_salaried=True,
        salary_employer="GLOBAL MEGA CORP",
        foir_pct=10,
        surplus_pct=40,
        has_bounces=False,
        variance_pct=0.0
    )
    
    # 2. Silver Moderate
    create_statement(
        filename=os.path.join(out_dir, 'scenario_silver_moderate.xlsx'),
        account_holder="Silver User",
        months_count=6,
        start_year=2025,
        start_month=1,
        base_salary=80000,
        is_salaried=True,
        salary_employer="MIDTIER SOLUTIONS",
        foir_pct=35,
        surplus_pct=20,
        has_bounces=False,
        variance_pct=0.0
    )
    
    # 3. Bronze Risky
    create_statement(
        filename=os.path.join(out_dir, 'scenario_bronze_risky.xlsx'),
        account_holder="Bronze User",
        months_count=4,
        start_year=2025,
        start_month=2,
        base_salary=45000,
        is_salaried=True,
        salary_employer="STARTUP INC",
        foir_pct=55,
        surplus_pct=5, # practically 0
        has_bounces=True,
        variance_pct=20.0 # 20% drops
    )
    
    # 4. Non Salaried
    create_statement(
        filename=os.path.join(out_dir, 'scenario_non_salaried.xlsx'),
        account_holder="Freelancer User",
        months_count=6,
        start_year=2025,
        start_month=1,
        base_salary=60000, # used for volume reference
        is_salaried=False,
        salary_employer="",
        foir_pct=15,
        surplus_pct=30,
        has_bounces=False,
        variance_pct=0.0
    )
    
    print("All datasets generated successfully in test_datasets/")
